# analysis/regression.py
import pandas as pd
import numpy as np
import statsmodels.api as sm
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import os

# ── 1. DATA ──────────────────────────────────────────────────────────────────
data = {
    'obs':   list(range(1, 25)),
    'sales': [504.72, 406.59, 398.55, 587.76, 598.92, 703.62, 387.24, 365.67,
              388.71, 372.96, 603.30, 614.73, 484.38, 227.76, 329.13, 308.25,
              433.86, 514.98, 404.70, 245.43, 433.20, 627.24, 647.61, 342.81],
    'prom':  [15.6, 22.2, 0.0, 0.0, 0.0, 31.8, 21.3, 3.9,
              0.0, 8.4, 45.3, 50.1, 39.6, 4.2, 0.0, 0.0,
              0.0, 13.8, 17.7, 0.0, 17.4, 37.8, 42.3, 11.4],
    'adv':   [30, 36, 45, 57, 39, 21, 12, 6,
              6, 30, 30, 33, 6, 33, 6, 3,
              45, 48, 0, 15, 9, 54, 36, 39],
    'index': [100, 102, 104, 104, 104, 100, 98, 96,
              98, 103, 105, 107, 107, 107, 108, 105,
              103, 108, 110, 112, 113, 112, 113, 114],
}
df = pd.DataFrame(data)

assert df.shape == (24, 5), f"Expected (24, 5), got {df.shape}"
assert abs(df['sales'].mean() - 455.505) < 0.01, f"Mean mismatch: {df['sales'].mean()}"
print("✓ Data loaded: 24 rows × 5 columns")
print(df.describe().round(2))

# ── 2. FEATURE ENGINEERING ───────────────────────────────────────────────────
df['prom_lag1'] = df['prom'].shift(1)
df['adv_lag1']  = df['adv'].shift(1)

# Seasonal dummies — assumption: Obs 1 = Q1 (Jan–Mar)
# Q4 is baseline (winter/cold) — omitted to avoid dummy trap
quarter = ((df['obs'] - 1) % 4) + 1
df['Q1'] = (quarter == 1).astype(int)
df['Q2'] = (quarter == 2).astype(int)
df['Q3'] = (quarter == 3).astype(int)

print("\n── Quarter distribution ──")
print(f"Q1 obs: {df['Q1'].sum()} (expect 6)")
print(f"Q2 obs: {df['Q2'].sum()} (expect 6)")
print(f"Q3 obs: {df['Q3'].sum()} (expect 6)")
print(f"Q4 obs: {(24 - df['Q1'].sum() - df['Q2'].sum() - df['Q3'].sum())} (expect 6)")
assert df['Q1'].sum() == df['Q2'].sum() == df['Q3'].sum() == 6
print("✓ Seasonal dummies correct")

print("\n── Correlation matrix (predictors) ──")
cols = ['prom', 'adv', 'index', 'prom_lag1', 'adv_lag1']
corr = df[cols].corr().round(2)
print(corr)
high_corr = [(c1, c2, corr.loc[c1, c2])
             for i, c1 in enumerate(cols)
             for j, c2 in enumerate(cols)
             if i < j and abs(corr.loc[c1, c2]) > 0.7]
if high_corr:
    print(f"\n⚠ High correlation pairs: {high_corr}")
else:
    print("\n✓ No severe multicollinearity detected")

# ── 3. MODELS ─────────────────────────────────────────────────────────────────
def run_ols(df_in, predictors, label):
    """Run OLS regression; returns fitted model."""
    d = df_in.dropna(subset=predictors + ['sales']).copy()
    X = sm.add_constant(d[predictors])
    y = d['sales']
    model = sm.OLS(y, X).fit()
    n = len(d)
    print(f"\n{'─'*60}")
    print(f"  {label}  (n={n})")
    print(f"{'─'*60}")
    print(f"  Adj. R²: {model.rsquared_adj:.4f}   R²: {model.rsquared:.4f}   AIC: {model.aic:.1f}")
    for var in model.params.index:
        sig = "***" if model.pvalues[var] < 0.01 else ("**" if model.pvalues[var] < 0.05 else ("*" if model.pvalues[var] < 0.10 else ""))
        print(f"  {var:15s}  coef={model.params[var]:8.3f}  t={model.tvalues[var]:6.2f}  p={model.pvalues[var]:.3f}  {sig}")
    return model

m1 = run_ols(df, ['prom', 'adv', 'index'], "M1: Base (prom + adv + index)")
m2 = run_ols(df, ['prom', 'adv', 'index', 'prom_lag1', 'adv_lag1'],
             "M2: Lags (+ prom_lag1 + adv_lag1)")
m3 = run_ols(df, ['prom', 'adv', 'index', 'Q1', 'Q2', 'Q3'],
             "M3: Seasonal (+ Q1 Q2 Q3 dummies)")
m4 = run_ols(df, ['prom', 'adv', 'index', 'prom_lag1', 'adv_lag1', 'Q1', 'Q2', 'Q3'],
             "M4: Full (lags + seasonal)")

print("\n── Model Comparison ──")
models = {'M1 Base': m1, 'M2 Lags': m2, 'M3 Seasonal': m3, 'M4 Full': m4}
comp = pd.DataFrame({
    name: {'Adj R²': m.rsquared_adj, 'R²': m.rsquared, 'AIC': m.aic, 'n': int(m.nobs)}
    for name, m in models.items()
}).T.round(4)
print(comp)

best_name = comp['Adj R²'].idxmax()
best_model = models[best_name]
best_predictors = [v for v in best_model.model.exog_names if v != 'const']
print(f"\n★ Best model: {best_name} (Adj R² = {comp.loc[best_name, 'Adj R²']:.4f})")

# ── 4. CASE QUESTIONS ─────────────────────────────────────────────────────────
print("\n" + "═"*60)
print("CASE QUESTION ANSWERS")
print("═"*60)

params = best_model.params
pvals  = best_model.pvalues

print("\n── Q1: $1,000 on Advertising or Promotion? ──")
prom_coef      = float(params.get('prom', 0) or 0)
adv_coef       = float(params.get('adv', 0) or 0)
prom_lag_coef  = float(params.get('prom_lag1', 0) or 0)
adv_lag_coef   = float(params.get('adv_lag1', 0) or 0)

prom_sig = pvals.get('prom', 1) < 0.05
adv_sig  = pvals.get('adv', 1)  < 0.05
prom_lag_sig = pvals.get('prom_lag1', 1) < 0.05
adv_lag_sig  = pvals.get('adv_lag1', 1)  < 0.05

print(f"  prom coefficient    : {prom_coef:.3f}  {'(significant)' if prom_sig else '(NOT significant)'}")
print(f"  adv  coefficient    : {adv_coef:.3f}  {'(significant)' if adv_sig else '(NOT significant)'}")
if prom_lag_coef is not None:
    print(f"  prom_lag1 coefficient: {prom_lag_coef:.3f}  {'(significant)' if prom_lag_sig else '(NOT significant)'}")
if adv_lag_coef is not None:
    print(f"  adv_lag1  coefficient: {adv_lag_coef:.3f}  {'(significant)' if adv_lag_sig else '(NOT significant)'}")

total_prom = (prom_coef or 0) + (prom_lag_coef or 0)
total_adv  = (adv_coef  or 0) + (adv_lag_coef  or 0)
print(f"\n  Total prom effect per $1K (current + lag): ${total_prom:.2f}K in sales")
print(f"  Total adv  effect per $1K (current + lag): ${total_adv:.2f}K in sales")
print("\n  → RECOMMENDATION: Spend on " + ("ADVERTISING" if total_adv >= total_prom else "PROMOTION"))

print("\n── Q2: Is Meat Loaf Mix Counter-Cyclical? ──")
index_coef = params.get('index', None)
index_sig  = pvals.get('index', 1) < 0.05
print(f"  index coefficient: {index_coef:.3f}  {'(significant p<0.05)' if index_sig else '(NOT significant)'}")
if index_coef < 0 and index_sig:
    print("  → CONFIRMED counter-cyclical: higher index (better economy) → lower sales")
elif index_coef > 0 and index_sig:
    print("  → NOT counter-cyclical (pro-cyclical): higher index → higher sales")
else:
    print("  → INCONCLUSIVE: economic index not a significant predictor in this model")

print("\n── Q3: Significant Seasonal Effects? ──")
seasonal_results = {}
for q in ['Q1', 'Q2', 'Q3']:
    if q in params:
        coef = params[q]
        pval = pvals[q]
        sig  = pval < 0.05
        seasonal_results[q] = {'coef': coef, 'pval': pval, 'sig': sig}
        print(f"  {q} vs Q4: coef={coef:.2f}  p={pval:.3f}  {'★ significant' if sig else 'not significant'}")

if seasonal_results and any(v['sig'] for v in seasonal_results.values()):
    print("\n  → CONFIRMED: Significant seasonal effects exist")
    best_q  = max(seasonal_results, key=lambda k: seasonal_results[k]['coef'])
    worst_q = min(seasonal_results, key=lambda k: seasonal_results[k]['coef'])
    print(f"  Highest sales quarter vs Q4 baseline: {best_q} (+{seasonal_results[best_q]['coef']:.1f}K)")
    print(f"  Lowest  sales quarter vs Q4 baseline: {worst_q} ({seasonal_results[worst_q]['coef']:.1f}K)")
else:
    print("\n  → NOT CONFIRMED: No seasonal dummies are statistically significant")

# ── 5. EXCEL OUTPUT ────────────────────────────────────────────────────────────
os.makedirs('output', exist_ok=True)
wb = Workbook()

header_font   = Font(bold=True, color="FFFFFF")
header_fill   = PatternFill("solid", fgColor="1F4E79")
sig_fill      = PatternFill("solid", fgColor="E2EFDA")
center        = Alignment(horizontal='center')
thin          = Side(style='thin')
border        = Border(left=thin, right=thin, top=thin, bottom=thin)

def style_header(cell):
    cell.font      = header_font
    cell.fill      = header_fill
    cell.alignment = center
    cell.border    = border

# Sheet 1: Data
ws_data = wb.active
ws_data.title = "Data"
headers = ['Obs', 'Sales ($K)', 'Prom ($K)', 'Adv ($K)', 'Index',
           'Prom Lag1', 'Adv Lag1', 'Q1', 'Q2', 'Q3']
for col, h in enumerate(headers, 1):
    c = ws_data.cell(row=1, column=col, value=h)
    style_header(c)
    ws_data.column_dimensions[get_column_letter(col)].width = 12

display_cols = ['obs', 'sales', 'prom', 'adv', 'index', 'prom_lag1', 'adv_lag1', 'Q1', 'Q2', 'Q3']
for r_idx, (_, row) in enumerate(df.iterrows(), 2):
    for c_idx, col in enumerate(display_cols, 1):
        val = row[col]
        ws_data.cell(row=r_idx, column=c_idx, value=round(float(val), 2) if pd.notna(val) else '')
        ws_data.cell(row=r_idx, column=c_idx).border = border

ws_data.cell(row=27, column=1, value='Mean').font = Font(bold=True)
for c_idx, col in enumerate(['sales', 'prom', 'adv', 'index'], 2):
    cell = ws_data.cell(row=27, column=c_idx)
    cell.value = round(df[col].mean(), 2)
    cell.font = Font(bold=True)
print("✓ Data sheet written")

# Sheet 2: Best Model Regression
ws_reg = wb.create_sheet("Regression Results")
ws_reg.merge_cells('A1:G1')
title_cell = ws_reg['A1']
title_cell.value     = f"OLS Regression Results — {best_name}"
title_cell.font      = Font(bold=True, size=14, color="FFFFFF")
title_cell.fill      = header_fill
title_cell.alignment = center

model_stats = [
    ('Dependent Variable:', 'Sales ($K)'),
    ('Observations:',       int(best_model.nobs)),
    ('R-squared:',          round(best_model.rsquared, 4)),
    ('Adjusted R-squared:', round(best_model.rsquared_adj, 4)),
    ('F-stat p-value:',     round(best_model.f_pvalue, 4)),
    ('AIC:',                round(best_model.aic, 2)),
]
for i, (label, val) in enumerate(model_stats, 3):
    ws_reg.cell(row=i, column=1, value=label).font = Font(bold=True)
    ws_reg.cell(row=i, column=2, value=val)

coef_headers = ['Variable', 'Coefficient', 'Std Error', 't-Statistic', 'P-value', '95% CI Low', '95% CI High']
for col, h in enumerate(coef_headers, 1):
    style_header(ws_reg.cell(row=10, column=col, value=h))
    ws_reg.column_dimensions[get_column_letter(col)].width = 16

conf = best_model.conf_int()
for r_idx, var in enumerate(best_model.params.index, 11):
    p = best_model.pvalues[var]
    row_data = [var, round(best_model.params[var], 4), round(best_model.bse[var], 4),
                round(best_model.tvalues[var], 4), round(p, 4),
                round(conf.loc[var, 0], 4), round(conf.loc[var, 1], 4)]
    for c_idx, val in enumerate(row_data, 1):
        cell = ws_reg.cell(row=r_idx, column=c_idx, value=val)
        cell.border = border
        if p < 0.05 and c_idx > 1:
            cell.fill = sig_fill
    ws_reg.cell(row=r_idx, column=1).font = Font(bold=(p < 0.05))

legend_row = 11 + len(best_model.params)
ws_reg.cell(row=legend_row + 1, column=1,
            value='Note: Green shading = statistically significant (p < 0.05)').font = Font(italic=True)
print("✓ Regression output sheet written")

# Sheet 3: Model Comparison
ws_comp = wb.create_sheet("Model Comparison")
ws_comp.merge_cells('A1:E1')
ws_comp['A1'].value     = 'Model Comparison — Selection by Adjusted R²'
ws_comp['A1'].font      = Font(bold=True, size=13, color="FFFFFF")
ws_comp['A1'].fill      = header_fill
ws_comp['A1'].alignment = center

for col, h in enumerate(['Model', 'Adj R²', 'R²', 'AIC', 'n'], 1):
    style_header(ws_comp.cell(row=3, column=col, value=h))
    ws_comp.column_dimensions[get_column_letter(col)].width = 16

for r_idx, (name, row) in enumerate(comp.iterrows(), 4):
    vals = [name, row['Adj R²'], row['R²'], row['AIC'], int(row['n'])]
    for c_idx, val in enumerate(vals, 1):
        ws_comp.cell(row=r_idx, column=c_idx, value=val).border = border
    if name == best_name:
        for c_idx in range(1, 6):
            ws_comp.cell(row=r_idx, column=c_idx).fill = sig_fill
            ws_comp.cell(row=r_idx, column=c_idx).font = Font(bold=True)

ws_comp.cell(row=r_idx + 2, column=1, value=f'★ Best model: {best_name}').font = Font(bold=True)
print("✓ Model comparison sheet written")

# Sheet 4: Case Question Answers
ws_qa = wb.create_sheet("Case Answers")
ws_qa.merge_cells('A1:C1')
ws_qa['A1'].value     = 'Magic Kitchens — Case Question Answers'
ws_qa['A1'].font      = Font(bold=True, size=13, color="FFFFFF")
ws_qa['A1'].fill      = header_fill
ws_qa['A1'].alignment = center
ws_qa.column_dimensions['A'].width = 30
ws_qa.column_dimensions['B'].width = 22
ws_qa.column_dimensions['C'].width = 55

def qa_row(ws, row_num, question, value, interpretation):
    ws.cell(row=row_num, column=1, value=question).font = Font(bold=True)
    ws.cell(row=row_num, column=2, value=value)
    cell = ws.cell(row=row_num, column=3, value=interpretation)
    cell.alignment = Alignment(wrap_text=True)
    ws.row_dimensions[row_num].height = 45

prom_total = (params.get('prom', 0) or 0) + (params.get('prom_lag1', 0) or 0)
adv_total  = (params.get('adv',  0) or 0) + (params.get('adv_lag1',  0) or 0)
q1_rec = "ADVERTISING" if adv_total >= prom_total else "PROMOTION"

qa_row(ws_qa, 3,
       'Q1: Prom or Adv ($1K)?',
       f"Recommend: {q1_rec}",
       f"$1K on promotion yields ~${prom_total:.2f}K in sales (current + lag). "
       f"$1K on advertising yields ~${adv_total:.2f}K in sales (current + lag).")

idx_val = params.get('index', 0)
idx_sig = pvals.get('index', 1) < 0.05
q2_answer = "CONFIRMED" if (idx_val < 0 and idx_sig) else "NOT CONFIRMED"
q2_interp = (f"Negative and significant index coefficient ({idx_val:.3f}, p<0.05): "
             "better economy → lower sales. Counter-cyclical confirmed.")  \
    if (idx_val < 0 and idx_sig) else \
    (f"Index coefficient ({idx_val:.3f}) is not statistically significant (p={pvals.get('index',1):.3f}). "
     "Cannot confirm counter-cyclical hypothesis from this data.")

qa_row(ws_qa, 5, 'Q2: Counter-cyclical?', q2_answer, q2_interp)

any_seas_sig = any(pvals.get(q, 1) < 0.05 for q in ['Q1', 'Q2', 'Q3'] if q in pvals)
q3_answer = "YES — Significant seasonal effects" if any_seas_sig else "NOT CONFIRMED"
q3_interp = ("Seasonal dummy coefficients are statistically significant. "
             "See regression sheet for quarter-by-quarter effects vs Q4 baseline.") \
    if any_seas_sig else \
    "No seasonal dummies reached p < 0.05. Seasonal effects are not statistically confirmed."

qa_row(ws_qa, 7, 'Q3: Seasonal effects?', q3_answer, q3_interp)
print("✓ Case Q&A sheet written")

# ── 6. FORECAST ───────────────────────────────────────────────────────────────
mean_prom  = df['prom'].mean()
mean_adv   = df['adv'].mean()
mean_index = df['index'].mean()
last_prom  = df['prom'].iloc[-1]
last_adv   = df['adv'].iloc[-1]
total_spend = mean_prom + mean_adv

scenario_profiles = {
    'Baseline': {'prom': mean_prom, 'adv': mean_adv},
    'Advertising-led': {'prom': round(total_spend * 0.30, 2), 'adv': round(total_spend * 0.70, 2)},
    'Promotion-led': {'prom': round(total_spend * 0.70, 2), 'adv': round(total_spend * 0.30, 2)},
}

forecast_rows = []
for scenario_name, spends in scenario_profiles.items():
    prev_prom = last_prom
    prev_adv = last_adv
    for i, q_num in enumerate([1, 2, 3, 4], start=1):
        row = {
            'scenario': scenario_name,
            'obs': 24 + len(forecast_rows) + 1,
            'quarter': f'Q{q_num}',
            'prom': spends['prom'],
            'adv': spends['adv'],
            'index': mean_index,
            'prom_lag1': prev_prom,
            'adv_lag1': prev_adv,
            'Q1': int(q_num == 1),
            'Q2': int(q_num == 2),
            'Q3': int(q_num == 3),
        }
        forecast_rows.append(row)
        prev_prom = spends['prom']
        prev_adv = spends['adv']

fc_df = pd.DataFrame(forecast_rows)
fc_X = sm.add_constant(fc_df[best_predictors], has_constant='add')
fc_df['forecast_sales'] = best_model.predict(fc_X)
fc_df['total_spend'] = fc_df['prom'] + fc_df['adv']

print("\n── Scenario Forecasts ──")
print(fc_df[['scenario', 'quarter', 'forecast_sales']].to_string(index=False))

# Forecast sheet (Sheet 5)
ws_fc = wb.create_sheet("Forecast")
ws_fc.merge_cells('A1:F1')
ws_fc['A1'].value     = 'Scenario Forecast (Constant Total Spend)'
ws_fc['A1'].font      = Font(bold=True, size=13, color="FFFFFF")
ws_fc['A1'].fill      = header_fill
ws_fc['A1'].alignment = center

for col, h in enumerate(['Scenario', 'Quarter', 'Forecast Sales ($K)', 'Prom ($K)', 'Adv ($K)', 'Total Spend ($K)'], 1):
    style_header(ws_fc.cell(row=3, column=col, value=h))
    ws_fc.column_dimensions[get_column_letter(col)].width = 20

for r_idx, (_, row) in enumerate(fc_df.iterrows(), 4):
    ws_fc.cell(row=r_idx, column=1, value=row['scenario'])
    ws_fc.cell(row=r_idx, column=2, value=row['quarter'])
    ws_fc.cell(row=r_idx, column=3, value=round(row['forecast_sales'], 2))
    ws_fc.cell(row=r_idx, column=4, value=round(row['prom'], 2))
    ws_fc.cell(row=r_idx, column=5, value=round(row['adv'], 2))
    ws_fc.cell(row=r_idx, column=6, value=round(row['total_spend'], 2))
    for c_idx in range(1, 7):
        ws_fc.cell(row=r_idx, column=c_idx).border = border

ws_fc.cell(row=17, column=1,
           value='Note: Baseline preserves the historical mean mix; alternative scenarios keep total spend constant and change only the mix.').font = Font(italic=True)
print("✓ Forecast sheet written")

# ── SAVE ──────────────────────────────────────────────────────────────────────
output_path = 'output/magic_kitchens_analysis.xlsx'
wb.save(output_path)
print(f"\n✓ Excel output saved → {output_path}")
print("  Sheets: Data | Regression Results | Model Comparison | Case Answers | Forecast")
